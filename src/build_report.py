import html
from pathlib import Path
from string import Template

import pandas as pd
import plotly.graph_objects as go
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR, SQL_RESULTS_DIR, READMISSION_COST, INTERVENTION_COST, RISK_REDUCTION
from src.db import query_df
from src.utils import load_metrics, save_metrics
from src import plot_style as ps

HEADER_FILL = PatternFill("solid", fgColor="1C5CAB")
TITLE_FONT = Font(bold=True, size=13, color="0B0B0B")
QUESTION_FONT = Font(bold=True, size=11, color="1C5CAB")


# excel
def style_sheet(ws):
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)


def write_block(ws, df, start_row, title=None):
    # writes a small table with a bold title and a coloured header row, returns next free row
    row = start_row
    if title:
        ws.cell(row=row, column=1, value=title).font = QUESTION_FONT
        row += 1
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=row, column=j, value=str(col))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
    for i, values in enumerate(df.itertuples(index=False), start=row + 1):
        for j, v in enumerate(values, start=1):
            ws.cell(row=i, column=j, value=None if pd.isna(v) else v)
    return row + len(df) + 3


def build_excel(m, findings, worklist):
    path = OUTPUT_DIR / "readmission_analysis_report.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    # summary
    ws = wb.create_sheet("Summary")
    ws.cell(row=1, column=1, value="Diabetes 30-Day Readmission Analysis - Executive Summary").font = TITLE_FONT
    kpis = pd.DataFrame(kpi_rows(m), columns=["Metric", "Value"])
    r = write_block(ws, kpis, 3, "Key numbers")
    r = write_block(ws, pd.DataFrame({"Finding": findings}), r, "Key findings")
    write_block(ws, pd.DataFrame({"Recommendation": recommendations(m)}), r, "Recommendations")
    style_sheet(ws)
    ws.column_dimensions["A"].width = 110

    # data quality
    ws = wb.create_sheet("Data Quality")
    r = write_block(ws, pd.read_csv(OUTPUT_DIR / "data_quality_checks.csv"), 1, "Checks run during cleaning")
    write_block(ws, pd.read_csv(OUTPUT_DIR / "data_quality_report.csv"), r, "Missing values per raw column")
    style_sheet(ws)

    # every sql result stacked on one sheet with its business question
    ws = wb.create_sheet("SQL Insights")
    index = pd.read_csv(SQL_RESULTS_DIR / "_query_index.csv")
    r = 1
    for q in index.itertuples():
        df = pd.read_csv(SQL_RESULTS_DIR / f"{q.query}.csv")
        r = write_block(ws, df, r, f"{q.query}: {q.question}")
    style_sheet(ws)

    ws = wb.create_sheet("Statistical Tests")
    r = write_block(ws, pd.read_csv(OUTPUT_DIR / "stats_chi_square_tests.csv"), 1,
                    "Chi-square tests on first stays (Bonferroni corrected)")
    r = write_block(ws, pd.read_csv(OUTPUT_DIR / "stats_mann_whitney_tests.csv"), r, "Mann-Whitney U tests")
    r = write_block(ws, pd.read_csv(OUTPUT_DIR / "stats_hba1c_ztest.csv"), r, "HbA1c two-proportion z-test")
    write_block(ws, pd.read_csv(OUTPUT_DIR / "stats_adjusted_odds_ratios.csv"), r,
                "Adjusted odds ratios (multivariable logistic regression)")
    style_sheet(ws)

    ws = wb.create_sheet("Model Performance")
    r = write_block(ws, pd.read_csv(OUTPUT_DIR / "model_metrics.csv"), 1, "Held-out test set, patient-level split")
    r = write_block(ws, pd.read_csv(OUTPUT_DIR / "model_risk_deciles.csv"), r, "Gradient boosting risk deciles")
    r = write_block(ws, pd.read_csv(OUTPUT_DIR / "model_feature_importance.csv"), r, "Permutation importance")
    write_block(ws, pd.read_csv(OUTPUT_DIR / "model_fairness_audit.csv"), r, "Fairness audit (top 20% threshold)")
    style_sheet(ws)

    ws = wb.create_sheet("Cost-Benefit")
    ws.cell(row=1, column=1, value=(f"Assumptions: ${READMISSION_COST:,} per readmission, ${INTERVENTION_COST} "
                                    f"per enrolled patient, {RISK_REDUCTION:.0%} risk reduction. Per 10,000 discharges."))
    r = write_block(ws, pd.read_csv(OUTPUT_DIR / "cost_strategy_comparison.csv"), 3, "Strategy comparison")
    write_block(ws, pd.read_csv(OUTPUT_DIR / "cost_sensitivity_grid.csv"), r, "Sensitivity grid")
    style_sheet(ws)

    ws = wb.create_sheet("Care Team Worklist")
    write_block(ws, worklist.head(500), 1, "Top 500 highest-risk stays (from vw_care_team_worklist in MySQL)")
    style_sheet(ws)

    wb.save(path)
    return path


# summary text
def money(x):
    # -1500 -> "-$1,500" instead of "$-1,500"
    return f"-${abs(x):,.0f}" if x < 0 else f"${x:,.0f}"


def kpi_rows(m):
    s, mo, c = m["sql"], m["model"], m["cost"]
    return [
        ("Eligible inpatient stays analysed", f"{s['stays']:,}"),
        ("Unique patients", f"{s['patients']:,}"),
        ("30-day readmission rate", f"{s['readmit_rate_pct']}%"),
        ("Readmissions (cost at ${:,} each)".format(READMISSION_COST), f"{s['readmissions']:,} (${s['readmission_cost_musd']:.0f}M)"),
        ("Stays with an HbA1c test", f"{s['hba1c_tested_pct']}%"),
        ("Model test ROC-AUC (95% CI)", f"{mo['hgb_test_auc']:.3f} ({mo['hgb_auc_ci'][0]:.3f}-{mo['hgb_auc_ci'][1]:.3f})"),
        ("LACE index test ROC-AUC", f"{mo['lace_test_auc']:.3f}"),
        ("Readmissions caught in top 20% (model vs LACE)", f"{mo['hgb_recall_top20_pct']}% vs {mo['lace_recall_top20_pct']}%"),
        ("Top risk decile readmission rate", f"{mo['hgb_top_decile_rate_pct']}% ({mo['hgb_lift_top10']}x average)"),
        ("Optimal enrolment", f"top {c['optimal_pct_targeted']:.0f}% of discharges"),
        ("Net savings per 10,000 discharges (optimal)", money(c["optimal_net_savings_per_10k"])),
        ("Net savings if everyone is enrolled", money(c["enrol_everyone_net_per_10k"])),
    ]


def build_findings(m):
    s, mo, c = m["sql"], m["model"], m["cost"]
    return [
        f"{s['readmit_rate_pct']}% of eligible diabetic stays ({s['readmissions']:,}) were readmitted within 30 days, "
        f"about ${s['readmission_cost_musd']:.0f}M at the AHRQ average cost.",
        f"Readmission risk climbs with every prior inpatient stay: {s['rate_0_prior']}% with none vs "
        f"{s['rate_5plus_prior']}% with 5+ ({s['rr_5plus_prior']}x). Adjusted OR per extra stay = {s['or_prior_inpatient']}.",
        f"{s['repeat_pct_patients']}% of patients (3+ stays) account for {s['repeat_pct_readmissions']}% of all readmissions.",
        f"Discharge to rehab / long-term care has a {s['rehab_rate']}% readmission rate vs {s['home_rate']}% for home "
        f"(adjusted OR {s['or_rehab']}). SNF discharges alone create ~{s['snf_excess']:,} excess readmissions.",
        f"Only {s['hba1c_tested_pct']}% of stays had an HbA1c test. For diabetes-primary first stays, tested patients "
        f"were readmitted {s['hba1c_tested_rate']}% vs {s['hba1c_untested_rate']}% untested (p = {s['hba1c_p']:.4f}).",
        f"The standard LACE index only reaches AUC {mo['lace_test_auc']:.3f}; the gradient boosting model reaches "
        f"{mo['hgb_test_auc']:.3f} on unseen patients (gain 95% CI {mo['auc_gain_vs_lace_ci'][0]:.3f} to "
        f"{mo['auc_gain_vs_lace_ci'][1]:.3f}).",
        f"Calling the top 20% of discharges catches {mo['hgb_recall_top20_pct']}% of readmissions with the model vs "
        f"{mo['lace_recall_top20_pct']}% with LACE - about {mo['hgb_recall_top20_pct'] / mo['lace_recall_top20_pct'] - 1:.0%} more with the same staff.",
        f"At ${INTERVENTION_COST}/patient and a {RISK_REDUCTION:.0%} effect, the break-even risk is {c['break_even_risk_pct']}%. "
        f"Enrolling the model's top {c['optimal_pct_targeted']:.0f}% nets ${c['optimal_net_savings_per_10k']:,.0f} per 10,000 "
        f"discharges; enrolling everyone loses ${-c['enrol_everyone_net_per_10k']:,.0f}.",
        f"The program stays profitable in {c['profitable_scenarios']} of {c['total_scenarios']} cost / effect scenarios tested.",
    ]


def recommendations(m):
    c = m["cost"]
    return [
        f"Run the risk model at discharge and enrol the top {c['optimal_pct_targeted']:.0f}-20% in a care transition "
        "program (follow-up call in 48h, med reconciliation, appointment booked before discharge).",
        "Build a 'frequent flyer' pathway for patients with 2+ admissions in the past year - under 9% of patients but over half of readmissions.",
        "Tighten hand-offs to SNF / rehab / LTC facilities (discharge summary + medication list sent same day).",
        "Make HbA1c testing standard for diabetic inpatients - currently only 1 in 6 stays is tested.",
        "Replace the LACE >= 10 rule with the model score, and re-check calibration and fairness every quarter.",
    ]


def collect_sql_numbers():
    # pull the headline numbers straight from the sql result files
    def q(name):
        return pd.read_csv(SQL_RESULTS_DIR / f"{name}.csv")

    k = q("q01_overall_kpis").iloc[0]
    prior = q("q03_prior_admissions_dose_response")
    disc = q("q04_discharge_destination").set_index("discharge_group")
    freq = q("q10_frequent_flyers").set_index("patient_type")
    ztest = pd.read_csv(OUTPUT_DIR / "stats_hba1c_ztest.csv").iloc[0]
    ors = pd.read_csv(OUTPUT_DIR / "stats_adjusted_odds_ratios.csv").set_index("term")["odds_ratio"]
    return {
        "stays": int(k["total_encounters"]),
        "patients": int(k["total_patients"]),
        "readmissions": int(k["readmissions_30d"]),
        "readmit_rate_pct": float(k["readmit_rate_pct"]),
        "readmission_cost_musd": round(k["readmission_cost_usd"] / 1e6, 1),
        "hba1c_tested_pct": float(k["hba1c_tested_pct"]),
        "rate_0_prior": float(prior.iloc[0]["readmit_rate_pct"]),
        "rate_5plus_prior": float(prior.iloc[-1]["readmit_rate_pct"]),
        "rr_5plus_prior": float(prior.iloc[-1]["relative_risk_vs_zero"]),
        "repeat_pct_patients": float(freq.loc["3+ stays", "pct_of_patients"]),
        "repeat_pct_readmissions": float(freq.loc["3+ stays", "pct_of_readmissions"]),
        "rehab_rate": float(disc.loc["Rehab / LTC", "readmit_rate_pct"]),
        "home_rate": float(disc.loc["Home", "readmit_rate_pct"]),
        "snf_excess": int(disc.loc["SNF / Nursing", "excess_readmissions_vs_home"]),
        "hba1c_tested_rate": float(ztest["rate_tested_pct"]),
        "hba1c_untested_rate": float(ztest["rate_not_tested_pct"]),
        "hba1c_p": float(ztest["p_value"]),
        "or_prior_inpatient": round(float(ors["prior_inpatient"]), 2),
        "or_rehab": round(float(ors["discharge_group = Rehab / LTC"]), 2),
    }


# html dashboard
def fig_layout(fig, title, xaxis, yaxis, height=380):
    fig.update_layout(
        title=dict(text=title, x=0, font=dict(size=15, color=ps.INK)),
        paper_bgcolor=ps.SURFACE, plot_bgcolor=ps.SURFACE, height=height,
        margin=dict(l=60, r=20, t=60, b=50),
        font=dict(family="system-ui, -apple-system, Segoe UI, sans-serif", color=ps.INK_2, size=12),
        xaxis=dict(title=xaxis, gridcolor=ps.GRID, linecolor=ps.AXIS, zeroline=False),
        yaxis=dict(title=yaxis, gridcolor=ps.GRID, linecolor=ps.AXIS, zeroline=False),
        legend=dict(orientation="h", y=-0.2), hoverlabel=dict(bgcolor="white"),
    )
    return fig


def table_html(df):
    return f"<details><summary>Show data table</summary>{df.to_html(index=False, border=0, classes='tbl')}</details>"


def build_dashboard(m, findings):
    charts = []

    prior = pd.read_csv(SQL_RESULTS_DIR / "q03_prior_admissions_dose_response.csv")
    fig = go.Figure(go.Bar(x=prior["prior_inpatient_visits"].astype(str), y=prior["readmit_rate_pct"],
                           marker_color=ps.BLUE, customdata=prior["encounters"],
                           hovertemplate="%{x} prior stays<br>%{y:.1f}% readmitted<br>%{customdata:,} stays<extra></extra>"))
    fig.update_xaxes(type="category")
    charts.append((fig_layout(fig, "Readmission rate by prior inpatient stays", "Inpatient stays in past year",
                              "Readmission rate (%)"), prior))

    disc = pd.read_csv(SQL_RESULTS_DIR / "q04_discharge_destination.csv").sort_values("readmit_rate_pct")
    fig = go.Figure(go.Bar(x=disc["readmit_rate_pct"], y=disc["discharge_group"], orientation="h",
                           marker_color=[ps.BLUE if r > m["sql"]["readmit_rate_pct"] else ps.LIGHT_GRAY
                                         for r in disc["readmit_rate_pct"]],
                           customdata=disc[["encounters", "excess_cost_usd"]],
                           hovertemplate="%{y}<br>%{x:.1f}% readmitted<br>%{customdata[0]:,} stays"
                                         "<br>excess cost vs home: $%{customdata[1]:,}<extra></extra>"))
    charts.append((fig_layout(fig, "Readmission rate by discharge destination", "Readmission rate (%)", ""), disc))

    dec = pd.read_csv(OUTPUT_DIR / "model_risk_deciles.csv")
    fig = go.Figure(go.Bar(x=dec["decile"], y=dec["readmit_rate_pct"],
                           marker_color=[ps.BLUE if d <= 2 else ps.LIGHT_GRAY for d in dec["decile"]],
                           customdata=dec[["lift", "cum_pct_of_readmissions"]],
                           hovertemplate="decile %{x}<br>%{y:.1f}% readmitted<br>lift %{customdata[0]}x"
                                         "<br>cumulative capture %{customdata[1]}%<extra></extra>"))
    fig.update_xaxes(dtick=1)
    charts.append((fig_layout(fig, "Actual readmission rate by model risk decile (1 = highest)", "Risk decile",
                              "Readmission rate (%)"), dec))

    curves = pd.read_csv(OUTPUT_DIR / "cost_benefit_curves.csv")
    fig = go.Figure()
    for name, color in [("Gradient boosting", ps.BLUE), ("LACE index", ps.ORANGE), ("No targeting (random)", ps.AXIS)]:
        c = curves[curves["strategy"] == name]
        fig.add_trace(go.Scatter(x=c["pct_targeted"], y=c["net_savings"], name=name, mode="lines",
                                 line=dict(color=color, width=2),
                                 hovertemplate=name + "<br>enrol top %{x}%<br>net $%{y:,.0f}<extra></extra>"))
    fig.update_layout(hovermode="x unified")
    charts.append((fig_layout(fig, "Net savings per 10,000 discharges vs % enrolled", "% of discharges enrolled",
                              "Net savings ($)"), curves[curves["pct_targeted"] % 5 == 0]))

    fair = pd.read_csv(OUTPUT_DIR / "model_fairness_audit.csv")
    fair_labels = (fair["attribute"].map({"race": "Race", "gender": "Gender", "age_band": "Age"}) + ": "
                   + fair["group"].replace({"AfricanAmerican": "African American"}))
    fig = go.Figure(go.Scatter(x=fair["auc"], y=fair_labels, mode="markers",
                               marker=dict(color=ps.BLUE, size=11, line=dict(color=ps.SURFACE, width=2)),
                               customdata=fair[["stays", "recall_pct"]],
                               hovertemplate="%{y}<br>AUC %{x:.3f}<br>n = %{customdata[0]:,}"
                                             "<br>catches %{customdata[1]}%<extra></extra>"))
    charts.append((fig_layout(fig, "Model AUC within each group (fairness check)", "AUC", "", height=420), fair))

    parts = []
    for i, (fig, data) in enumerate(charts):
        chart_div = fig.to_html(full_html=False, include_plotlyjs="cdn" if i == 0 else False,
                                config={"displayModeBar": False})
        parts.append(f"<section class='card'>{chart_div}{table_html(data)}</section>")

    mo, c, s = m["model"], m["cost"], m["sql"]
    tiles = [
        ("Readmission rate", f"{s['readmit_rate_pct']}%", f"{s['readmissions']:,} of {s['stays']:,} stays"),
        ("Cost of readmissions", f"${s['readmission_cost_musd']:.0f}M", f"at ${READMISSION_COST:,} each"),
        ("HbA1c tested", f"{s['hba1c_tested_pct']}%", "of diabetic inpatient stays"),
        ("Model AUC", f"{mo['hgb_test_auc']:.3f}", f"vs LACE {mo['lace_test_auc']:.3f}"),
        ("Caught in top 20%", f"{mo['hgb_recall_top20_pct']}%", f"vs {mo['lace_recall_top20_pct']}% with LACE"),
        ("Net savings", f"${c['optimal_net_savings_per_10k'] / 1000:,.0f}K", "per 10,000 discharges"),
    ]
    tile_html = "".join(f"<div class='tile'><div class='label'>{html.escape(a)}</div><div class='value'>{html.escape(b)}</div>"
                        f"<div class='sub'>{html.escape(d)}</div></div>" for a, b, d in tiles)
    finding_html = "".join(f"<li>{html.escape(f)}</li>" for f in findings)

    # page layout lives in dashboard_template.html, we just fill in the pieces
    template = Template((Path(__file__).parent / "dashboard_template.html").read_text(encoding="utf-8"))
    page = template.substitute(tiles=tile_html, findings=finding_html, charts="".join(parts))
    path = OUTPUT_DIR / "readmission_dashboard.html"
    path.write_text(page, encoding="utf-8")
    return path


def main():
    ps.set_style()
    save_metrics("sql", collect_sql_numbers())
    m = load_metrics()
    findings = build_findings(m)

    # the care team list comes straight out of the mysql view
    worklist = query_df("SELECT * FROM vw_care_team_worklist WHERE risk_tier = 'High' ORDER BY predicted_risk_pct DESC")
    worklist.to_csv(OUTPUT_DIR / "care_team_worklist.csv", index=False)

    xlsx = build_excel(m, findings, worklist)
    dash = build_dashboard(m, findings)

    print(f"report  -> {xlsx.name}")
    print(f"dashboard -> {dash.name}")
    print(f"worklist -> care_team_worklist.csv ({len(worklist):,} high-risk stays)")
    print("\nKEY FINDINGS")
    for f in findings:
        print(" -", f)


if __name__ == "__main__":
    main()
